from app.agents.document_agent import _requested_slide_count
from app.agents.supervisor import _forced_agent_specs


def test_explicit_mission_routes_downloadable_presentation_to_document_agent():
    query = (
        "Create a downloadable three-slide PowerPoint about CRAG vs Self-RAG. "
        "Slide 1: goals. Slide 2: side-by-side retrieval flow. Slide 3: validation checks."
    )

    assert _forced_agent_specs(query) == [
        {"agent": "document", "query": query, "doc_type": "pptx"}
    ]


def test_explicit_mission_preserves_research_and_artifact_specialists():
    query = "Research current hybrid retrieval practice and create a PowerPoint deck."

    assert [task["agent"] for task in _forced_agent_specs(query)] == [
        "document",
        "web_research",
    ]


def test_step_based_mission_does_not_silently_return_to_ordinary_chat():
    query = (
        "Create a three-step plan to compare CRAG and Self-RAG, explain each, "
        "then provide a comparison and final verification checklist."
    )

    assert [task["agent"] for task in _forced_agent_specs(query)] == ["web_research"]


def test_document_agent_honors_numeric_and_worded_total_slide_counts():
    assert _requested_slide_count("Create a three-slide PowerPoint") == 3
    assert _requested_slide_count("Generate a 7 slide deck") == 7
    assert _requested_slide_count("Make a one-slide deck") == 1
    assert _requested_slide_count("Create a general presentation") is None


def test_pptx_renderer_preserves_requested_total_slide_count():
    from io import BytesIO
    from pptx import Presentation

    from app.agents.document_agent import _build_pptx

    slides = [
        {"title": f"Section {number}", "bullets": [f"Detail {number}"]}
        for number in range(1, 12)
    ]
    presentation = Presentation(BytesIO(_build_pptx("Audit deck", slides)))

    # The agent's count includes its dedicated title slide.
    assert len(presentation.slides) == 12
    assert presentation.slides[0].shapes.title.text == "Audit deck"
    assert presentation.slides[-1].shapes.title.text == "Section 11"


async def test_document_agent_repairs_short_plan_before_export(monkeypatch, tmp_path):
    from unittest.mock import AsyncMock

    from pptx import Presentation
    from app.agents.artifacts import ArtifactStore
    from app.agents import document_agent

    short_plan = {
        "title": "Audit deck",
        "slides": [{"title": "Section 1", "bullets": ["Detail 1"]}],
    }
    complete_plan = {
        "title": "Audit deck",
        "slides": [
            {"title": f"Section {number}", "bullets": [f"Detail {number}"]}
            for number in range(1, 4)
        ],
    }
    llm = AsyncMock(side_effect=[short_plan, complete_plan])
    artifact_store = ArtifactStore(tmp_path / "artifacts")
    monkeypatch.setattr(document_agent, "call_llm_json", llm)
    monkeypatch.setattr(document_agent, "artifact_store", artifact_store)

    result = await document_agent.DocumentAgent().run({
        "query": "Create a four-slide presentation about reliability",
        "user_id": "test-user",
        "session_id": "test-session",
    })

    assert result.success is True
    assert result.metadata["slide_count"] == 4
    assert llm.await_count == 2
    assert result.artifacts[0]["type"] == "pptx"
    artifact_path = artifact_store.get_path(result.artifacts[0]["id"])
    assert artifact_path is not None and artifact_path.is_file()
    exported = Presentation(artifact_path)
    assert len(exported.slides) == 4


async def test_document_agent_does_not_export_short_deck_after_failed_repair(monkeypatch):
    from types import SimpleNamespace
    from unittest.mock import AsyncMock

    from app.agents import document_agent

    short_plan = {"title": "Audit deck", "slides": [{"title": "Only content slide", "bullets": []}]}
    llm = AsyncMock(side_effect=[short_plan, short_plan])
    create_artifact = AsyncMock(return_value=SimpleNamespace(to_dict=lambda: {}))
    monkeypatch.setattr(document_agent, "call_llm_json", llm)
    monkeypatch.setattr(document_agent.artifact_store, "create", create_artifact)

    result = await document_agent.DocumentAgent().run({
        "query": "Create a four-slide presentation about reliability",
        "user_id": "test-user",
        "session_id": "test-session",
    })

    assert result.success is False
    assert "contained only 2 after one repair attempt" in result.error
    create_artifact.assert_not_awaited()
